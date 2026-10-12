from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "量化投资多因子研发平台_项目计划书.docx"

FONT_CN = "Noto Sans CJK SC"
FONT_EN = "Arial"
NAVY = "243B53"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F7F9"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT_CN
    rfonts = run._element.get_or_add_rPr().rFonts
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(key), FONT_CN)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_run_font(run, 9, color=RGBColor(80, 80, 80))


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        normal._element.rPr.rFonts.set(qn(key), FONT_CN)
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.24

    title = doc.styles["Title"]
    title.font.name = FONT_CN
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        title._element.rPr.rFonts.set(qn(key), FONT_CN)
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 16, 14, 7),
        ("Heading 2", 13, 11, 5),
        ("Heading 3", 11, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = FONT_CN
        for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            style._element.rPr.rFonts.set(qn(key), FONT_CN)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "量化投资多因子研发平台    项目计划书"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in hp.runs:
        set_run_font(run, 8.5, color=RGBColor(90, 90, 90))

    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)


def add_para(doc, text="", bold_lead=None, align=None, keep=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if keep:
        p.paragraph_format.keep_with_next = True
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1.5)
        run = p.add_run(item)
        set_run_font(run)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(1.5)
        run = p.add_run(f"{index}. {item}")
        set_run_font(run)


def add_table(doc, headers, rows, widths=None, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(text)
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                set_run_font(run, font_size, True, RGBColor(255, 255, 255))
        if widths:
            cell.width = Inches(widths[i])
    for r_idx, row_data in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cell = cells[i]
            cell.text = str(value)
            set_cell_border(cell)
            set_cell_margins(cell)
            if r_idx % 2 == 1:
                set_cell_shading(cell, PALE_GRAY)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Inches(widths[i])
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) <= 14 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    set_run_font(run, font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, bold=True, color=BLACK)
    return p


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("量化投资多因子研发平台")
    set_run_font(run, 27, True, BLACK)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(10)
    run = p2.add_run("项目计划书")
    set_run_font(run, 18, True, BLACK)
    doc.add_paragraph()
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p3.add_run("课程项目初版报告")
    set_run_font(run, 11, color=RGBColor(80, 80, 80))
    for _ in range(5):
        doc.add_paragraph()
    rows = [
        ("课程名称", "待填写"),
        ("指导教师", "待填写"),
        ("项目组长", "待填写"),
        ("项目成员", "待填写"),
        ("计划周期", "2026年10月至2026年12月"),
        ("文档版本", "V1.0"),
    ]
    add_table(doc, ["项目", "内容"], rows, widths=[1.8, 4.3], font_size=10)


def build():
    doc = Document()
    style_document(doc)
    add_cover(doc)

    summary_heading = add_heading(doc, "项目摘要", 1)
    summary_heading.paragraph_format.page_break_before = True
    add_para(doc, "本项目建设一个面向少量研究人员的量化投资多因子研发平台。平台部分负责用户权限、数据版本、任务调度、日志和结果归档；产品部分负责因子配置、单因子检验、图表和研究报告。首期限定沪深A股日频数据，以课程验收为目标，在2026年12月中旬前完成可运行、可复现、可演示的研究闭环。")
    add_para(doc, "项目采用模块化单体和独立异步Worker。Web、API、Worker、PostgreSQL与Redis通过Docker Compose统一部署，行情和因子矩阵采用Parquet保存并由DuckDB查询。该方案保留后续拆分空间，同时控制四人团队在十周左右周期内的实现风险。")

    add_heading(doc, "项目结论", 2)
    add_bullets(doc, [
        "首期完成数据导入、因子选择、异步计算、统计检验、图表展示、报告生成和实验复现。",
        "六个内置因子覆盖动量、反转、波动率、换手率、市值和市盈率。",
        "四位成员各承担一个占责任25%的验收包，技术工作量允许不同，但成果必须独立可核验。",
        "实时行情、自动交易、机器学习和复杂微服务明确排除在首期范围之外。",
    ])

    add_heading(doc, "目录", 1)
    add_numbered(doc, [
        "项目背景与目标",
        "项目范围",
        "用户和业务流程",
        "系统架构",
        "数据与研究口径",
        "开发计划",
        "团队分工",
        "测试与验收",
        "风险管理",
        "预期成果和待确认事项",
    ])
    first_section = add_heading(doc, "一 项目背景与目标", 1)
    first_section.paragraph_format.page_break_before = True
    add_heading(doc, "项目背景", 2)
    add_para(doc, "因子研究涉及行情获取、数据清洗、因子计算、统计检验、任务执行和报告编制。若成员各自使用零散脚本，数据口径、参数和输出难以统一，也无法稳定复现历史实验。本项目用一个统一平台固化这些共性能力，并把研究方法组织为可操作的业务功能。")
    add_heading(doc, "项目目标", 2)
    add_numbered(doc, [
        "支持管理员导入、校验、发布和追踪数据版本。",
        "支持研究员选择内置因子、样本、日期和预处理参数并提交后台任务。",
        "统一计算IC、Rank IC、ICIR、分组收益、多空收益、换手率和相关性。",
        "自动生成结构化结果和HTML研究报告，并记录完整复现信息。",
        "在单台Windows电脑上完成私有部署，固化数据后核心流程可以离线运行。",
    ])
    add_heading(doc, "验收指标", 2)
    add_table(doc, ["指标", "验收标准"], [
        ("流程完整性", "登录至报告下载全流程运行成功"),
        ("单因子性能", "300只股票和3年日频数据在60秒内完成"),
        ("页面响应", "普通查询在推荐环境中不超过2秒"),
        ("计算正确性", "平台结果与黄金数据独立验算一致"),
        ("结果复现", "报告可追溯数据版本、因子版本、参数和代码版本"),
        ("演示时间", "核心业务演示控制在10分钟以内"),
    ], widths=[1.6, 4.8])

    add_heading(doc, "二 项目范围", 1)
    add_heading(doc, "首期功能", 2)
    add_table(doc, ["模块", "首期内容"], [
        ("身份权限", "管理员、研究员、只读用户三种角色"),
        ("数据管理", "CSV和Parquet导入、校验、版本、预览和统计"),
        ("因子管理", "六个内置因子、参数说明和不可变版本"),
        ("研究任务", "后台执行、进度、日志、失败原因和手动重试"),
        ("因子检验", "预处理、IC系列指标、分组收益、换手率和相关性"),
        ("结果报告", "结果图表、HTML报告、下载和复现信息"),
        ("私域部署", "Docker Compose统一启动，核心流程离线运行"),
    ], widths=[1.45, 4.95])
    add_heading(doc, "优先级", 2)
    add_table(doc, ["等级", "内容"], [
        ("P0 必须完成", "权限、数据版本、内置因子、异步任务、单因子检验、图表、报告和Docker部署"),
        ("P1 应当完成", "历史指数成分、简单多因子合成、任务取消和PDF导出"),
        ("P2 时间允许", "Ollama报告摘要、表达式因子编辑器和扩展监控"),
    ], widths=[1.5, 4.9])
    add_heading(doc, "非目标", 2)
    add_para(doc, "首期不实现实时行情、分钟或Tick数据、自动交易、券商接口、资金账户、复杂多因子优化、机器学习预测、网页任意Python执行、移动端、Kubernetes、多机容灾和生产级高并发。")

    add_heading(doc, "三 用户和业务流程", 1)
    add_heading(doc, "用户角色", 2)
    add_table(doc, ["角色", "主要权限"], [
        ("管理员", "管理用户、数据集、数据版本和全部研究任务"),
        ("研究员", "查看授权数据，创建任务，查看和下载自己的结果"),
        ("只读用户", "查看授权结果和报告，不得发起或修改任务"),
    ], widths=[1.5, 4.9])
    add_heading(doc, "业务流程", 2)
    add_numbered(doc, [
        "管理员导入数据并完成字段、日期、重复值和价格关系校验。",
        "管理员发布数据版本，研究员选择已发布数据和因子版本。",
        "研究员设置股票样本、日期范围和预处理参数并提交任务。",
        "Worker执行数据读取、预处理、因子计算和检验，页面显示进度。",
        "任务成功后保存指标、图表数据和报告，用户查看复现信息并下载成果。",
    ])
    add_heading(doc, "核心演示", 2)
    add_para(doc, "管理员查看数据版本；研究员选择20日动量因子和沪深300样本，提交异步任务并查看进度；任务完成后打开IC、分组收益和换手率图表，下载报告；最后切换只读账号验证权限限制。")

    add_heading(doc, "四 系统架构", 1)
    add_heading(doc, "架构选择", 2)
    add_para(doc, "首期采用模块化单体和独立异步Worker。FastAPI内部按身份权限、数据、因子、研究、任务和报告划分模块。React仅访问FastAPI，不能直接连接数据库或文件存储。该结构降低部署和联调成本，又为以后拆出计算、数据和报告服务保留接口边界。")
    add_table(doc, ["组件", "职责", "技术"], [
        ("Web", "研究工作台和结果展示", "React TypeScript"),
        ("API", "权限、业务规则和接口", "FastAPI"),
        ("Worker", "长耗时因子计算和报告生成", "Celery"),
        ("元数据库", "用户、版本、任务、指标摘要", "PostgreSQL"),
        ("队列缓存", "任务消息、进度和短期缓存", "Redis"),
        ("分析存储", "行情、因子矩阵和计算结果", "Parquet DuckDB"),
        ("部署", "跨Mac和Windows统一环境", "Docker Compose"),
    ], widths=[1.15, 3.15, 2.1])
    add_heading(doc, "私域部署", 2)
    add_para(doc, "首次拉取镜像和采集Tushare数据需要网络。数据快照准备完成后，登录、计算、检验和报告均在本机运行，业务数据不上传外部服务。本地模型只作为P2组件草拟摘要，关闭后不影响任何核心功能。")

    add_heading(doc, "五 数据与研究口径", 1)
    add_heading(doc, "数据范围", 2)
    add_para(doc, "演示数据采用沪深300相关固定样本和约3至5年日频记录，同时准备20至50只股票的小型集成测试集。固定样本可能产生幸存者偏差，报告必须明确披露；历史指数成分作为P1功能。")
    add_table(doc, ["数据表", "主要用途"], [
        ("证券主表", "股票代码、市场、上市日和退市日"),
        ("交易日历", "交易日判断和日期对齐"),
        ("日行情", "开高低收、成交量和成交额"),
        ("复权因子", "生成统一收益口径"),
        ("每日指标", "市值、换手率和估值因子"),
        ("证券状态", "停牌、ST和交易限制"),
        ("行业历史", "行业中性化"),
        ("基准行情", "计算基准和超额表现"),
    ], widths=[1.55, 4.85])
    add_heading(doc, "因子和检验", 2)
    add_table(doc, ["类型", "首期内容"], [
        ("内置因子", "20日动量、5日反转、20日波动率、20日换手率、市值和市盈率"),
        ("预处理", "去极值、标准化、可选市值和行业中性化"),
        ("统计检验", "覆盖率、IC、Rank IC、IC均值、IC标准差和ICIR"),
        ("收益检验", "分组收益、多空收益、换手率和因子相关性"),
    ], widths=[1.55, 4.85])
    add_heading(doc, "时间和复权规则", 2)
    add_bullets(doc, [
        "系统保存未复权行情和复权因子，研究收益采用后复权口径。",
        "因子值在交易日收盘后形成，最早用于下一交易日收益检验。",
        "财务数据按实际披露日生效，不按报告期提前使用。",
        "黄金数据由独立Excel或Python脚本验算，平台结果必须在约定容差内一致。",
    ])

    add_heading(doc, "六 开发计划", 1)
    add_table(doc, ["阶段", "日期", "主要成果"], [
        ("需求与计划", "10月9日至11日", "PRD、计划书、架构、分工和学习路线"),
        ("原型与骨架", "10月12日至25日", "数据样例、页面原型、工程骨架和Docker环境"),
        ("核心平台", "10月26日至11月15日", "数据、权限、任务调度和因子计算"),
        ("研究产品", "11月16日至29日", "检验指标、结果页面和HTML报告"),
        ("联调部署", "11月30日至12月6日", "联调、测试、私域部署和可选加分项"),
        ("验收答辩", "12月7日至13日", "修复、个人验收包、汇报和彩排"),
    ], widths=[1.3, 1.65, 3.45], font_size=8.5)
    add_heading(doc, "三次汇报", 2)
    add_numbered(doc, [
        "第一次汇报说明目标、范围、架构、原型、数据样例和团队分工。",
        "第二次汇报展示数据版本、权限、异步任务和单因子计算链路。",
        "第三次汇报运行完整演示，展示结果、报告、复现信息和个人验收材料。",
    ])

    add_heading(doc, "七 团队分工", 1)
    add_para(doc, "四位成员的责任占比均为25%。责任比例按可验收成果确定，不按代码行数或编程难度平均。每人必须提交署名材料，并能够独立回答自己模块的设计和验收问题。")
    add_table(doc, ["成员", "责任包", "主要任务", "验收材料", "占比"], [
        ("A 队长", "架构与集成", "后端、任务、集成、部署", "架构、接口、部署和测试记录", "25%"),
        ("B", "数据工程", "数据、字典、清洗和质量", "字段字典、样例和质量报告", "25%"),
        ("C", "因子研究", "公式、验算和报告内容", "因子说明、黄金结果和差异记录", "25%"),
        ("D", "产品测试", "原型、测试、手册和展示", "原型、用例、缺陷和演示材料", "25%"),
    ], widths=[0.75, 1.15, 1.7, 2.2, 0.6], font_size=8.2)
    add_heading(doc, "协作规则", 2)
    add_bullets(doc, [
        "队长统一管理代码仓库和正式提交，其他成员按约定目录交付材料。",
        "个人材料记录作者、版本、日期、输入、产出、自测结论和复核人。",
        "每周至少一次内部同步，逐项更新状态、风险、下周任务和验收人。",
        "P0需求冻结后，新增需求只能进入P1或P2。",
    ])

    add_heading(doc, "八 测试与验收", 1)
    add_table(doc, ["层级", "数据或范围", "负责人"], [
        ("单元测试", "5只股票和30个交易日的人工构造数据", "A与C"),
        ("集成测试", "20至50只股票的固定Parquet快照", "A与B"),
        ("正确性测试", "Excel或独立Python脚本形成黄金结果", "C"),
        ("端到端测试", "登录、提交任务、结果和报告下载", "D"),
        ("部署测试", "Windows主演示环境和Mac开发环境", "A与D"),
    ], widths=[1.35, 3.85, 1.2])
    add_heading(doc, "个人验收", 2)
    add_para(doc, "项目建立四份个人验收模板。A提交架构与集成材料，B提交数据质量材料，C提交公式和独立验算材料，D提交原型、测试和展示材料。导师汇报前完成交叉复核，最终由队长归档。")

    add_heading(doc, "九 风险管理", 1)
    add_table(doc, ["风险", "影响", "应对措施"], [
        ("编程工作集中", "队长成为进度瓶颈", "冻结P0，其他成员完成独立验收责任包"),
        ("量化知识不足", "公式和时间对齐错误", "共同学习、口径评审和黄金数据验算"),
        ("接口或网络异常", "数据采集或演示失败", "提前固化Parquet数据，演示离线运行"),
        ("幸存者偏差", "研究表现被高估", "报告披露限制，历史成分列为P1"),
        ("范围膨胀", "核心闭环延期", "新增需求进入P1或P2，不影响P0"),
        ("跨平台差异", "Mac与Windows行为不同", "Docker、WSL2、统一路径和持续测试"),
    ], widths=[1.35, 2.1, 2.95], font_size=8.5)

    add_heading(doc, "十 预期成果和待确认事项", 1)
    add_heading(doc, "预期成果", 2)
    add_bullets(doc, [
        "可在Windows环境运行的量化因子研发平台。",
        "完整源代码、Docker配置、数据库迁移和测试代码。",
        "固定样例数据、数据字典和质量报告。",
        "六个内置因子及单因子检验结果。",
        "HTML研究报告、使用说明、答辩材料和演示脚本。",
        "PRD、团队计划、学习路线和四份个人验收包。",
    ])
    add_heading(doc, "待确认事项", 2)
    add_table(doc, ["事项", "当前状态"], [
        ("课程名称和指导教师", "待填写"),
        ("四名成员姓名和学号", "待填写"),
        ("三次导师汇报日期", "待导师确认"),
        ("每名成员每周投入时间", "待团队确认"),
        ("学校文档和答辩格式", "待确认"),
    ], widths=[3.5, 2.9])
    doc.core_properties.title = "量化投资多因子研发平台项目计划书"
    doc.core_properties.subject = "课程项目开发计划和小组分工"
    doc.core_properties.author = "项目团队"
    doc.core_properties.keywords = "量化投资 多因子 平台 项目计划"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
